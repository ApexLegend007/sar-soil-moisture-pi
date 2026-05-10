"""
Export all experiment outputs (metrics + plots) to Excel workbooks.

One workbook per experiment folder under output/.
Each workbook has:
  - A "Metrics" sheet with the flattened JSON metrics as a table.
  - One sheet per PNG plot, with the image embedded and its filename as the sheet name.

Run:
    uv run python experiments/classification_new_data/code/export_to_excel.py
"""

import json
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from constants import OUTPUT_PATH


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ALT_FILL    = PatternFill("solid", fgColor="D6E4F0")
THIN        = Side(style="thin", color="AAAAAA")
BORDER      = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _style_header(cell):
    cell.fill   = HEADER_FILL
    cell.font   = HEADER_FONT
    cell.border = BORDER
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _style_data(cell, alt=False):
    cell.fill   = ALT_FILL if alt else PatternFill("solid", fgColor="FFFFFF")
    cell.border = BORDER
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _auto_width(ws):
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 60)


# ---------------------------------------------------------------------------
# JSON → flat rows
# ---------------------------------------------------------------------------

def _flatten(obj, prefix=""):
    """Recursively flatten a nested dict into (key, value) pairs."""
    rows = []
    if isinstance(obj, dict):
        for k, v in obj.items():
            label = f"{prefix} > {k}" if prefix else str(k)
            if isinstance(v, dict):
                rows.extend(_flatten(v, label))
            else:
                rows.append((label, v))
    else:
        rows.append((prefix, obj))
    return rows


def _load_all_metrics(folder: Path):
    """Return a list of (source_file, [(key, value), ...]) for every JSON in folder."""
    results = []
    for jf in sorted(folder.glob("*.json")):
        try:
            data = json.loads(jf.read_text())
        except Exception:
            continue
        rows = _flatten(data)
        results.append((jf.name, rows))
    return results


# ---------------------------------------------------------------------------
# Sheet builders
# ---------------------------------------------------------------------------

def _write_metrics_sheet(wb, all_metrics):
    ws = wb.create_sheet("Metrics")
    ws.freeze_panes = "A2"

    headers = ["Source File", "Metric", "Value"]
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        _style_header(cell)

    row_num = 2
    for source_file, rows in all_metrics:
        for ki, (key, val) in enumerate(rows):
            alt = (row_num % 2 == 0)
            ws.cell(row=row_num, column=1, value=source_file).fill = ALT_FILL if alt else PatternFill("solid", fgColor="FFFFFF")
            ws.cell(row=row_num, column=1).border = BORDER
            ws.cell(row=row_num, column=1).alignment = Alignment(vertical="center")

            c2 = ws.cell(row=row_num, column=2, value=key)
            _style_data(c2, alt)

            raw = val
            if isinstance(val, float):
                raw = round(val, 6)
            c3 = ws.cell(row=row_num, column=3, value=raw)
            _style_data(c3, alt)

            row_num += 1

    _auto_width(ws)


def _write_plot_sheet(wb, png_path: Path):
    # Sheet names max 31 chars, no special chars
    sheet_name = png_path.stem[:31]
    # Deduplicate sheet names
    existing = [s.title for s in wb.worksheets]
    base, n = sheet_name, 1
    while sheet_name in existing:
        sheet_name = f"{base[:28]}_{n}"
        n += 1

    ws = wb.create_sheet(sheet_name)
    ws.cell(row=1, column=1, value=png_path.name).font = Font(bold=True, size=12)

    try:
        img = XLImage(str(png_path))
        # Scale down large images to fit nicely
        max_px = 900
        if img.width > max_px:
            scale = max_px / img.width
            img.width  = int(img.width  * scale)
            img.height = int(img.height * scale)
        ws.add_image(img, "A3")
    except Exception as e:
        ws.cell(row=3, column=1, value=f"[Could not embed image: {e}]")


# ---------------------------------------------------------------------------
# Per-folder workbook
# ---------------------------------------------------------------------------

def export_folder(folder: Path):
    all_metrics = _load_all_metrics(folder)

    # Collect plots — top-level PNGs and any plots/ subdirectory
    png_files = sorted(folder.glob("*.png")) + sorted((folder / "plots").glob("*.png") if (folder / "plots").exists() else [])

    if not all_metrics and not png_files:
        return  # nothing to export

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default empty sheet

    if all_metrics:
        _write_metrics_sheet(wb, all_metrics)

    for png in png_files:
        _write_plot_sheet(wb, png)

    out_path = folder / f"{folder.name}_results.xlsx"
    wb.save(out_path)
    print(f"  Saved → {out_path.relative_to(OUTPUT_PATH.parent.parent.parent)}")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    print(f"Scanning {OUTPUT_PATH} ...\n")
    folders = sorted(p for p in OUTPUT_PATH.iterdir() if p.is_dir())
    if not folders:
        print("No output folders found.")
        return

    for folder in folders:
        print(f"[{folder.name}]")
        export_folder(folder)

    print("\nDone.")


if __name__ == "__main__":
    main()
